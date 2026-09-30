#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""独立复核：重新打开产物 xlsx，按官方规则逐条验，不信任生成器的自校验。"""

import re
import sys
from pathlib import Path

from openpyxl import load_workbook

D = Path("/Users/zmk/Documents/vsc/golang/orderhub/outputs/wecom")
IMPORT_XLSX = D / "企业微信通讯录导入-成员信息.xlsx"
REF_XLSX = D / "企业微信组织架构与导入说明.xlsx"

EXPECT_HEADERS = ["帐号", "姓名", "别名", "职务", "部门", "性别", "手机", "座机", "个人邮箱"]
ACCOUNT_RE = re.compile(r"^[A-Za-z0-9._-]{1,32}$")
MOBILE_RE = re.compile(r"^(\+[0-9]{6,15}|1[3-9][0-9]{9})$")
EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$")

ok, bad = [], []


def check(cond, msg):
    (ok if cond else bad).append(msg)


# ---------- 导入表 ----------
wb = load_workbook(IMPORT_XLSX)
check(wb.sheetnames == ["成员信息"], f"导入表只含一个工作表「成员信息」，实际 {wb.sheetnames}")
ws = wb["成员信息"]
headers = [c.value for c in ws[1]]
check(headers == EXPECT_HEADERS, f"表头与官方字段口径一致：{headers}")

rows = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if all(v in (None, "") for v in r):
        continue
    rows.append(dict(zip(headers, r)))
check(len(rows) == 71, f"成员共 {len(rows)} 人")
check(len(rows) <= 200, f"成员数 {len(rows)} ≤ 未认证企业上限 200")

accts, mails, mobiles = {}, {}, {}
max_len = 0
for i, r in enumerate(rows, start=2):
    tag = f"第{i}行 {r['姓名']}"
    check(bool(ACCOUNT_RE.match(str(r["帐号"]))), f"{tag} 帐号合法 {r['帐号']}")
    check(bool(str(r["姓名"]).strip() == str(r["姓名"]) and str(r["姓名"])), f"{tag} 姓名非空无首尾空格")
    check(bool(r["部门"]), f"{tag} 部门非空")
    check(str(r["部门"]).startswith("taskpilot怀民/"), f"{tag} 部门路径以企业名开头")
    check(len(str(r["部门"])) <= 32, f"{tag} 部门路径 {len(r['部门'])} 字符 ≤ 32")
    max_len = max(max_len, len(str(r["部门"])))
    check(r["性别"] in ("男", "女"), f"{tag} 性别取值 {r['性别']}")
    check(bool(r["个人邮箱"]) or bool(r["手机"]), f"{tag} 手机/个人邮箱二选一已满足")
    if r["个人邮箱"]:
        check(bool(EMAIL_RE.match(str(r["个人邮箱"]))), f"{tag} 邮箱格式合法")
    if r["手机"]:
        check(bool(MOBILE_RE.match(str(r["手机"]))), f"{tag} 手机格式合法")
    for col in EXPECT_HEADERS:
        check("," not in str(r[col] or "") and "，" not in str(r[col] or ""),
              f"{tag} 字段「{col}」不含逗号")
        check(str(r[col] or "") == str(r[col] or "").strip(),
              f"{tag} 字段「{col}」无首尾空格")
    accts[r["帐号"]] = accts.get(r["帐号"], 0) + 1
    if r["个人邮箱"]:
        mails[r["个人邮箱"]] = mails.get(r["个人邮箱"], 0) + 1
    if r["手机"]:
        mobiles[r["手机"]] = mobiles.get(r["手机"], 0) + 1

check(all(v == 1 for v in accts.values()), f"帐号全局唯一（{len(accts)} 个不重复）")
check(all(v == 1 for v in mails.values()), f"个人邮箱全局唯一（{len(mails)} 个不重复）")
check(all(v == 1 for v in mobiles.values()), "手机号全局唯一")

# ---------- 与说明表交叉核对 ----------
wb2 = load_workbook(REF_XLSX)
check(wb2.sheetnames == ["组织架构总览", "成员花名册", "部门负责人", "校验报告", "迭代日志", "导入说明"],
      f"说明表工作表：{wb2.sheetnames}")

ws_org = wb2["组织架构总览"]
declared_depts = {}
for r in ws_org.iter_rows(min_row=2, values_only=True):
    if r[2]:
        declared_depts[r[2]] = r[3]
check(len(declared_depts) == 31, f"声明部门 {len(declared_depts)} 个")

used = sorted({r["部门"] for r in rows})
missing = [p for p in used if p not in declared_depts]
check(not missing, f"导入表中出现的 {len(used)} 条部门路径全部在组织架构表中有定义（缺失 {missing}）")

ws_mem = wb2["成员花名册"]
ref_rows = [r for r in ws_mem.iter_rows(min_row=2, values_only=True) if r[1]]
check(len(ref_rows) == len(rows), f"花名册人数 {len(ref_rows)} 与导入表 {len(rows)} 一致")
check(sorted(r[1] for r in ref_rows) == sorted(r["姓名"] for r in rows), "花名册姓名集合与导入表一致")

ws_rep = wb2["校验报告"]
report = [r for r in ws_rep.iter_rows(min_row=2, values_only=True) if r[0] and r[0] != "—"]
check(not report, f"校验报告无遗留问题（{len(report)} 条）")

ws_log = wb2["迭代日志"]
log_rows = [r for r in ws_log.iter_rows(min_row=2, values_only=True) if r[0]]
check(len(log_rows) >= 2, f"迭代日志 {len(log_rows)} 轮")
check(log_rows[0][1] > log_rows[-1][1], f"错误数随迭代下降：{log_rows[0][1]} -> {log_rows[-1][1]}")
check(log_rows[-1][1] == 0, "最终一轮错误数为 0")

# ---------- 汇总 ----------
DEPTH = {}
for p in used:
    DEPTH[p] = p.count("/") + 1
print("=" * 78)
print("独立复核结果")
print("=" * 78)
print(f"成员 {len(rows)} 人 / 部门 {len(declared_depts)} 个 / 使用到的部门路径 {len(used)} 条")
print(f"最深部门层级 {max(DEPTH.values())} 层（上限 15）/ 最长部门路径 {max_len} 字符（上限 32）")
print(f"通过 {len(ok)} 项，失败 {len(bad)} 项")
if bad:
    print("-" * 78)
    for b in bad:
        print("  ✗", b)
print("-" * 78)
print("前 5 行预览：")
for r in rows[:5]:
    print(f"  {r['帐号']:<16} {r['姓名']:<6} {r['职务']:<14} {r['部门']}")
sys.exit(1 if bad else 0)
